"""Excel 导出（openpyxl）。"""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .scraper import Product

HEADERS = [
    ("排名", 8),
    ("产品名称", 28),
    ("描述", 55),
    ("关键词", 40),
    ("分类", 22),
    ("ProductHunt原始链接", 50),
    ("跳转后真实链接", 50),
    ("票数", 8),
    ("评论数", 8),
    ("日排名", 8),
    ("周排名", 8),
    ("月排名", 8),
    ("发布时间", 26),
    ("缩略图", 50),
]


def export_excel(products: list[Product], path: str, sheet_title: str = "ProductHunt"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="DA552F")  # PH 橙色
    for col, (name, width) in enumerate(HEADERS, 1):
        c = ws.cell(row=1, column=col, value=name)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"

    # 按票数排序（榜单顺序）
    products = sorted(products, key=lambda p: -p.votes)
    for i, p in enumerate(products, 1):
        row = [
            i, p.name, p.tagline, p.keywords, p.category,
            p.ph_url, p.real_url, p.votes, p.comments,
            p.daily_rank, p.weekly_rank, p.monthly_rank,
            p.featured_at, p.thumbnail,
        ]
        for col, val in enumerate(row, 1):
            c = ws.cell(row=i + 1, column=col, value=val)
            c.alignment = Alignment(vertical="top", wrap_text=col in (3, 4))
        for col in (6, 7, 14):
            v = ws.cell(row=i + 1, column=col).value
            if v:
                ws.cell(row=i + 1, column=col).hyperlink = v
                ws.cell(row=i + 1, column=col).font = Font(color="0563C1", underline="single")

    wb.save(path)
    return path


def _rows(products: list[Product]):
    products = sorted(products, key=lambda p: -p.votes)
    for i, p in enumerate(products, 1):
        yield [i, p.name, p.tagline, p.keywords, p.category,
               p.ph_url, p.real_url, p.votes, p.comments,
               p.daily_rank, p.weekly_rank, p.monthly_rank,
               p.featured_at, p.thumbnail]


def export_csv(products: list[Product], path: str, **_):
    import csv
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([h for h, _ in HEADERS])
        w.writerows(_rows(products))
    return path


def export_json(products: list[Product], path: str, **_):
    import json
    keys = [h for h, _ in HEADERS]
    data = [dict(zip(keys, row)) for row in _rows(products)]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    return path


EXPORTERS = {".xlsx": export_excel, ".csv": export_csv, ".json": export_json}


def export_any(products: list[Product], path: str, sheet_title: str = "ProductHunt"):
    """按扩展名自动选择导出格式（xlsx/csv/json）。"""
    from pathlib import Path
    fn = EXPORTERS.get(Path(path).suffix.lower(), export_excel)
    return fn(products, path, sheet_title=sheet_title)


def default_filename(period: str, label: str) -> str:
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"producthunt_{period}_{label}_{ts}.xlsx"
